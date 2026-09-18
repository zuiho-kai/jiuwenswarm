"""Review must use the current artifact without exposing benchmark answers."""
import asyncio
import json

from jiuwenswarm.benchmarks.duplex_review_datasets import ReviewEnvironment, tree_hash


def test_office_agents_only_receive_initial_workbooks(tmp_path):
    assets = tmp_path / "assets"
    data = assets / "spreadsheetbench_verified_400"
    source = data / "spreadsheet/13-1"
    source.mkdir(parents=True)
    (source / "1_13-1_init.xlsx").write_bytes(b"original workbook")
    (source / "1_13-1_golden.xlsx").write_bytes(b"hidden answer")
    (data / "dataset.json").write_text(json.dumps([
        {"id": "13-1", "spreadsheet_path": "spreadsheet/13-1", "instruction": "Merge rows."}]))
    env = ReviewEnvironment("office", tmp_path / "run", assets)
    for directory in env.directories.values():
        assert [p.name for p in (directory / "inputs").iterdir()] == ["1_13-1_init.xlsx"]
    assert asyncio.run(env.review_snapshot())["available"] is False
    owner = env.directories["coordinator"] / "final"
    owner.mkdir()
    (owner / "1_13-1_output.xlsx").write_bytes(b"candidate")
    result = asyncio.run(env.review_snapshot())
    assert result["patch_sha256"] == tree_hash(owner)
    assert (env.directories["reviewer"] / "final/1_13-1_output.xlsx").read_bytes() == b"candidate"


def test_javascript_review_hash_tracks_changes_and_preserves_review_tests(tmp_path):
    env = ReviewEnvironment.__new__(ReviewEnvironment)
    env.suite, env.owner, env.lock = "development", "implementer", asyncio.Lock()
    env.directories = {name: tmp_path / name for name in ("implementer", "reviewer")}
    source = env.directories["implementer"] / "src"
    source.mkdir(parents=True)
    (source / "index.js").write_text("export default 'old'")
    reviewer = env.directories["reviewer"]
    (reviewer / "src").mkdir(parents=True)
    (reviewer / "src/removed.js").write_text("old snapshot")
    (reviewer / "regression.test.js").write_text("independent checks")
    before = asyncio.run(env.review_snapshot())
    (source / "index.js").write_text("export default 'fixed'")
    after = asyncio.run(env.review_snapshot())
    assert before["patch_sha256"] != after["patch_sha256"]
    assert not (reviewer / "src/removed.js").exists()
    assert (reviewer / "regression.test.js").read_text() == "independent checks"
    assert (reviewer / "src/index.js").read_text() == "export default 'fixed'"


def test_snapshot_waits_for_current_write_to_commit(tmp_path):
    async def scenario():
        env = ReviewEnvironment.__new__(ReviewEnvironment)
        env.suite, env.owner, env.lock = "office", "coordinator", asyncio.Lock()
        env.directories = {name: tmp_path / name for name in ("coordinator", "reviewer")}
        source = env.directories["coordinator"] / "final"
        source.mkdir(parents=True)
        workbook = source / "output.xlsx"
        workbook.write_bytes(b"partial")
        async with env.lock:
            pending = asyncio.create_task(env.review_snapshot())
            await asyncio.sleep(0)
            assert not pending.done()
            workbook.write_bytes(b"committed")
        await pending
        assert (env.directories["reviewer"] / "final/output.xlsx").read_bytes() == b"committed"
    asyncio.run(scenario())


def test_live_evidence_contains_actual_change_and_failure_without_grader(tmp_path):
    async def scenario():
        env = ReviewEnvironment.__new__(ReviewEnvironment)
        env.suite, env.owner, env.lock, env.task = "development", "implementer", asyncio.Lock(), "Use Redis"
        env.directories = {"implementer": tmp_path / "implementer"}
        source = env.directories["implementer"] / "src"
        source.mkdir(parents=True)
        file = source / "queue.js"
        file.write_text("queue = 'Kafka'\n")
        env._review_files = env._artifact_files()
        file.write_text("queue = 'Redis'\n")
        evidence = await env.review_evidence("run queue test", "exit_code=1\nconnection refused")
        assert evidence["changed_files"] == ["queue.js"]
        assert "+queue = 'Redis'" in evidence["artifact_changes"]
        assert "connection refused" in evidence["tool_output"]
        repeated = await env.review_evidence("read queue", "exit_code=0")
        assert repeated["revision"] == evidence["revision"]
        assert not repeated["changed_files"]
    asyncio.run(scenario())


def test_live_evidence_skips_unchanged_success_and_baseline_repro(tmp_path):
    async def scenario():
        env = ReviewEnvironment.__new__(ReviewEnvironment)
        env.suite, env.owner, env.lock, env.task = "development", "implementer", asyncio.Lock(), "Fix issue"
        env.directories = {"implementer": tmp_path / "implementer"}
        source = env.directories["implementer"] / "src"
        source.mkdir(parents=True)
        (source / "index.js").write_text("old")
        env._review_files = env._artifact_files()
        env._baseline_files = env._review_files.copy()
        clean = await env.review_evidence("ls -la", "exit_code=0\nindex.js")
        assert clean["review_reason"] == "skip"
        for label in ("Expected:", "Expected value to be:"):
            repro = await env.review_evidence("jest issue", f"exit_code=1\nTests: 1 failed\n{label} x\nReceived: y")
            assert repro["review_reason"] == "skip"
        broken_command = await env.review_evidence("node src", "exit_code=1\nSyntaxError: Unexpected token")
        assert broken_command["review_reason"] == "tool_failure"
        miss = await env.review_evidence('grep -n "Utils.s" src/utils.js', "exit_code=1\n")
        assert miss["review_reason"] == "skip"
        assert miss["skip_reason"] == "search returned no matches"
        missing_file = await env.review_evidence("grep pattern missing.js", "exit_code=2\ngrep: missing.js: No such file")
        assert missing_file["review_reason"] == "tool_failure"
        compound = await env.review_evidence("grep pattern src/index.js; false", "exit_code=1\n")
        assert compound["review_reason"] == "tool_failure"
        (source / "index.js").write_text("bad fix")
        changed = await env.review_evidence("edit", "exit_code=0\n")
        assert changed["review_reason"] == "artifact_change"
        candidate = await env.review_evidence("jest issue", "exit_code=1\nTests: 1 failed\nExpected: x\nReceived: y")
        assert candidate["review_reason"] == "tool_failure"
        (source / "index.js").write_text("good fix")
        assert not await env.review_is_current(changed)
    asyncio.run(scenario())


def test_successful_analysis_is_reviewed_in_bounded_batches(tmp_path):
    async def scenario():
        env = ReviewEnvironment.__new__(ReviewEnvironment)
        env.suite, env.owner, env.lock, env.task = "office", "coordinator", asyncio.Lock(), "Group input rows"
        env.directories = {"coordinator": tmp_path / "coordinator"}
        env._review_files = {}
        env._input_context = [{"sheet": "input", "sample_rows": [[1, "batch", 20]]}]
        first = await env.review_evidence("python3 -c 'inspect()'", "exit_code=0\n1 row")
        assert first["review_reason"] == "skip"
        second = await env.review_evidence("python3 -c 'filter_rows()'", "exit_code=0\nData rows: 0")
        assert second["review_reason"] == "analysis"
        assert len(second["analysis_batch"]) == 2
        assert "Data rows: 0" in second["analysis_batch"][-1]["output"]
        assert "batch" in second["input_context"]
        for _ in range(4):
            await env.review_evidence("python3 inspect.py", "exit_code=0\nanalysis")
        assert env._analysis_reviews == 3
        exhausted = await env.review_evidence("python3 inspect.py", "exit_code=0\nanalysis")
        assert exhausted["review_reason"] == "skip"
        clean = await env.review_evidence("ls -la", "exit_code=0\ninputs")
        assert clean["review_reason"] == "skip"
    asyncio.run(scenario())


def test_preview_contains_original_input_only_and_does_not_modify_it(tmp_path):
    import hashlib
    import openpyxl

    assets = tmp_path / "assets"
    data = assets / "spreadsheetbench_verified_400"
    source = data / "spreadsheet/13-1"
    source.mkdir(parents=True)
    original = source / "1_13-1_init.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["public input marker", 23])
    workbook.save(original)
    workbook.close()
    (source / "1_13-1_golden.xlsx").write_text("SECRET_GRADING_ANSWER")
    (data / "dataset.json").write_text(json.dumps([
        {"id": "13-1", "spreadsheet_path": "spreadsheet/13-1", "instruction": "Group rows."}]))
    before = hashlib.sha256(original.read_bytes()).hexdigest()
    env = ReviewEnvironment("office", tmp_path / "run", assets)
    env.enable_live_review()
    assert "public input marker" in env.task
    assert "SECRET_GRADING_ANSWER" not in env.task
    assert hashlib.sha256(original.read_bytes()).hexdigest() == before


def test_submission_checks_use_only_public_input_and_current_artifact(tmp_path):
    import openpyxl

    env = ReviewEnvironment.__new__(ReviewEnvironment)
    env.suite, env.owner, env.lock = "office", "coordinator", asyncio.Lock()
    env.directories = {"coordinator": tmp_path}
    (tmp_path / "inputs").mkdir()
    (tmp_path / "final").mkdir()
    wb = openpyxl.Workbook()
    wb.active.title = "Input"
    wb.create_sheet("Output")
    wb.save(tmp_path / "inputs/task_init.xlsx")
    assert not asyncio.run(env.submission_state())["ready"]
    wb.save(tmp_path / "final/task_output.xlsx")
    assert asyncio.run(env.submission_state())["ready"]
    del wb["Input"]
    wb.save(tmp_path / "final/task_output.xlsx")
    assert not asyncio.run(env.submission_state())["ready"]
    wb.close()
